"""Text out of a document a founder uploads: PDF, Word, PowerPoint, and plain
text (Markdown, CSV, JSON). What comes out is only ever read as information;
it is treated as untrusted input like any web page.

Parsing is CPU work: call `extract_text` off the event loop.
"""

import io
import json
import zipfile
from pathlib import PurePath

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError as DocxPackageError
from pptx import Presentation
from pptx.exc import PackageNotFoundError as PptxPackageError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

# Extensions, so a browser's guess at the content type doesn't decide it.
PDF, DOCX, PPTX = ".pdf", ".docx", ".pptx"
PLAIN = (".txt", ".md", ".markdown", ".csv", ".tsv", ".json")
SUPPORTED = (PDF, DOCX, PPTX, *PLAIN)


class UnsupportedDocument(ValueError):
    """Not a kind of file we can read."""


class UnreadableDocument(ValueError):
    """A supported kind, but no text could be got out of it."""


def extension(filename: str) -> str:
    return PurePath(filename).suffix.lower()


def _pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise UnreadableDocument("This PDF is password-protected.")
        pages = [page.extract_text() or "" for page in reader.pages]
    except PdfReadError as exc:
        raise UnreadableDocument("This PDF looks damaged.") from exc
    return "\n\n".join(f"[Page {i}]\n{text.strip()}" for i, text in enumerate(pages, 1) if text.strip())


def _docx(data: bytes) -> str:
    document = DocxDocument(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _pptx(data: bytes) -> str:
    deck = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(deck.slides, 1):
        texts = [shape.text_frame.text.strip() for shape in slide.shapes
                 if shape.has_text_frame and shape.text_frame.text.strip()]
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        if texts or notes:
            slides.append(f"[Slide {i}]\n" + "\n".join(texts) + (f"\nNotes: {notes}" if notes else ""))
    return "\n\n".join(slides)


def _plain(data: bytes, ext: str) -> str:
    text = data.decode("utf-8", errors="replace")
    if ext == ".json":
        try:
            return json.dumps(json.loads(text), indent=1, ensure_ascii=False)
        except json.JSONDecodeError:
            return text
    return text


def extract_text(filename: str, data: bytes) -> str:
    """The document's text, pages or slides marked. Raises
    UnsupportedDocument for a kind we don't read and UnreadableDocument when
    there is no text in it (a scanned PDF, a deck of images)."""
    ext = extension(filename)
    if ext not in SUPPORTED:
        raise UnsupportedDocument(f"I can't read {ext or 'that kind of'} files yet. Try PDF, Word, PowerPoint or text.")
    try:
        if ext == PDF:
            text = _pdf(data)
        elif ext == DOCX:
            text = _docx(data)
        elif ext == PPTX:
            text = _pptx(data)
        else:
            text = _plain(data, ext)
    except UnreadableDocument:
        raise
    # What the libraries raise for a broken or mislabelled file.
    except (ValueError, KeyError, OSError, zipfile.BadZipFile, DocxPackageError, PptxPackageError) as exc:
        raise UnreadableDocument(f"{filename} couldn't be opened; it may be damaged.") from exc
    if not text.strip():
        raise UnreadableDocument(f"There's no text in {filename}: if it's a scan or images, send a text version.")
    return text.strip()
