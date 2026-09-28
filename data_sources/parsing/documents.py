from pathlib import Path

from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation

from utils.log import logger


def parse_document(file_path: str) -> str:
    """Parse a document file and return clean text content."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    suffix = path.suffix.lower()
    logger.info("parse_document: parsing %s (%s)", path.name, suffix)

    match suffix:
        case ".pdf":
            return _parse_pdf(path)
        case ".docx":
            return _parse_docx(path)
        case ".pptx":
            return _parse_pptx(path)
        case ".txt" | ".md":
            return path.read_text(encoding="utf-8")
        case _:
            raise ValueError(f"Unsupported file type: {suffix}")


def _parse_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text:
            pages.append(f"--- Page {i + 1} ---\n{text}")
    return "\n\n".join(pages)


def _parse_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)
    return "\n\n".join(paragraphs)


def _parse_pptx(path: Path) -> str:
    prs = Presentation(str(path))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        texts.append(text)
        if texts:
            slides.append(f"--- Slide {i} ---\n" + "\n".join(texts))
    return "\n\n".join(slides)
