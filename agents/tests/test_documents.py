"""Text out of the documents founders share: what each format gives, and the
files we refuse with a reason the founder can act on."""

import io

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from data_sources.documents import UnreadableDocument, UnsupportedDocument, extract_text


def _docx() -> bytes:
    doc = Document()
    doc.add_paragraph("Shanga sells fuel telemetry to truck fleets.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Plan", "KES 1,500 per truck"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _pptx() -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Traction"
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = "40 fleets in Kenya"
    slide.notes_slide.notes_text_frame.text = "Mention the Twiga pilot"
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def test_word_text_and_tables():
    text = extract_text("one-pager.docx", _docx())
    assert "fuel telemetry" in text and "Plan | KES 1,500 per truck" in text


def test_slides_are_numbered_with_their_notes():
    text = extract_text("Deck.PPTX", _pptx())
    assert text.startswith("[Slide 1]") and "40 fleets in Kenya" in text and "Notes: Mention the Twiga pilot" in text


def test_plain_text_formats():
    assert extract_text("customers.csv", b"name,city\nTwiga,Nairobi") == "name,city\nTwiga,Nairobi"
    assert '"plan": "pro"' in extract_text("pricing.json", b'{"plan":"pro"}')


def test_unsupported_damaged_and_empty_files_say_why():
    with pytest.raises(UnsupportedDocument, match="Try PDF"):
        extract_text("setup.exe", b"MZ")
    with pytest.raises(UnreadableDocument, match="damaged"):
        extract_text("deck.pptx", b"not really a deck")
    with pytest.raises(UnreadableDocument, match="no text"):
        extract_text("blank.txt", b"   \n")
